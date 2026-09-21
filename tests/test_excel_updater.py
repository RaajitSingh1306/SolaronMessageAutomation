import os
import sys
import tempfile

import openpyxl
import pytest

# Add the project root to sys.path so we can import services
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from services.excel.updater import (  # noqa: E402
    _find_col_index,
    update_excel_status,
)
from config import Config


@pytest.fixture
def empty_sheet():
    wb = openpyxl.Workbook()
    return wb.active



def test_find_col_index_success(empty_sheet):
    """Test that it correctly finds the column index based on keywords."""
    header_row = ["ID", "Plant Name", "Daily Energy", "Message Status"]
    empty_sheet.append(header_row)

    # "plant", "name" should match "Plant Name" which is column 2
    idx = _find_col_index(empty_sheet, 1, "plant", "name")
    assert idx == 2

    # "message status" should match "Message Status" which is column 4
    idx = _find_col_index(empty_sheet, 1, "message status")
    assert idx == 4


def test_find_col_index_not_found(empty_sheet):
    """Test that it returns -1 if column is not found."""
    header_row = ["ID", "Unknown", "Data"]
    empty_sheet.append(header_row)

    idx = _find_col_index(empty_sheet, 1, "plant", "name")
    assert idx == -1


def test_update_excel_status_empty_plants():
    """Test that it returns True immediately if sent_plants is empty."""
    assert update_excel_status("fake_path.xlsx", []) is True


def test_update_excel_status_file_not_found():
    """Test that it returns False if the file doesn't exist."""
    assert update_excel_status("does_not_exist.xlsx", [("Plant A", 2)]) is False


def test_update_excel_status_success():
    """Test the full success path of marking plants as done."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Report"

    # Header at row Config.EXCEL_DATA_START_ROW - 1
    Config.EXCEL_DATA_START_ROW = 2
    
    # Row 1: Header
    sheet.append(["ID", "Plant Name", "Message Status"])
    # Row 2: Plant A
    sheet.append([1, "Gopal Naidu", ""])
    # Row 3: Plant B
    sheet.append([2, "Mandar Keskar", ""])
    # Row 4: Plant C
    sheet.append([3, "Vivek Mandir", ""])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        filepath = tmp.name

    try:
        wb.save(filepath)
        wb.close()

        # Test updating Plant A and Plant C
        sent_plants = [("Gopal Naidu", 2), ("vivek mandir", 4)]  # Test case insensitivity
        result = update_excel_status(filepath, sent_plants)
        assert result is True

        # Verify changes
        wb_updated = openpyxl.load_workbook(filepath)
        sheet_updated = wb_updated["Report"]

        assert sheet_updated.cell(row=2, column=3).value == "done"
        assert sheet_updated.cell(row=3, column=3).value in ("", None)
        assert sheet_updated.cell(row=4, column=3).value == "done"

        # Verify backup was created
        assert os.path.exists(filepath + ".backup")
    finally:
        try:
            wb_updated.close()
        except Exception:
            pass
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except Exception:
                pass
        if os.path.exists(filepath + ".backup"):
            try:
                os.remove(filepath + ".backup")
            except Exception:
                pass


def test_update_excel_status_fallback_columns():
    """Test fallback when 'Plant Name' or 'Message Status' are missing from header."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Report"

    # Header at row Config.EXCEL_DATA_START_ROW - 1
    Config.EXCEL_DATA_START_ROW = 2

    # Row 1: Header missing Plant Name and Message Status!
    # According to logic, plant defaults to column 2, status gets appended at max_column + 1
    sheet.append(["ID", "Some Other Data", "More Data"])

    # Because plant defaults to column 2, we put plant name there
    sheet.append([1, "Gopal Naidu", "Data"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        filepath = tmp.name

    try:
        wb.save(filepath)
        wb.close()

        sent_plants = [("Gopal Naidu", 2)]
        result = update_excel_status(filepath, sent_plants)
        assert result is True

        wb_updated = openpyxl.load_workbook(filepath)
        sheet_updated = wb_updated["Report"]

        # Original max_column was 3, so Message Status should be at column 4
        assert sheet_updated.cell(row=1, column=4).value == "Message Status"
        assert sheet_updated.cell(row=2, column=4).value == "done"
        wb_updated.close()
    finally:
        try:
            wb_updated.close()
        except Exception:
            pass
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except Exception:
                pass
        if os.path.exists(filepath + ".backup"):
            try:
                os.remove(filepath + ".backup")
            except Exception:
                pass


def test_update_excel_status_fallback_row_0():
    """Test that if row_index is 0, it falls back to name scan."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Report"

    # Header at row Config.EXCEL_DATA_START_ROW - 1
    Config.EXCEL_DATA_START_ROW = 2

    # Row 1: Header
    sheet.append(["ID", "Plant Name", "Message Status"])
    # Row 2: Plant A
    sheet.append([1, "Gopal Naidu", ""])
    # Row 3: Plant B
    sheet.append([2, "Mandar Keskar", ""])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        filepath = tmp.name

    try:
        wb.save(filepath)
        wb.close()

        # Test updating Plant B with row_index=0
        sent_plants = [("Mandar Keskar", 0)]
        result = update_excel_status(filepath, sent_plants)
        assert result is True

        wb_updated = openpyxl.load_workbook(filepath)
        sheet_updated = wb_updated["Report"]

        assert sheet_updated.cell(row=2, column=3).value in ("", None)
        assert sheet_updated.cell(row=3, column=3).value == "done"
        wb_updated.close()
    finally:
        try:
            wb_updated.close()
        except Exception:
            pass
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except Exception:
                pass
        if os.path.exists(filepath + ".backup"):
            try:
                os.remove(filepath + ".backup")
            except Exception:
                pass
